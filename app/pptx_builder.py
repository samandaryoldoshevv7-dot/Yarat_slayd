"""Shablon (.pptx) asosida tayyor taqdimot yig'ish.

Shablondagi namuna slaydlar o'chiriladi, dizayn (master/layout) saqlanib qoladi va
AI matni shu layoutlar asosida yangi slaydlarga joylanadi.
"""
import copy
import io
import logging
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

log = logging.getLogger(__name__)

TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
BODY_TYPES = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.SUBTITLE}
LABELS = {
    "uz": {"agenda": "Reja", "thanks": "E'tiboringiz uchun rahmat!"},
    "ru": {"agenda": "План", "thanks": "Спасибо за внимание!"},
    "en": {"agenda": "Agenda", "thanks": "Thank you for your attention!"},
}


def _ph_types(layout) -> list:
    return [ph.placeholder_format.type for ph in layout.placeholders]


def _has_title(layout) -> bool:
    return any(t in TITLE_TYPES for t in _ph_types(layout))


def _has_body(layout) -> bool:
    return any(t in BODY_TYPES for t in _ph_types(layout))


def _title_on_top(layout) -> bool:
    """Sarlavha matn maydonidan yuqorida joylashgan oddiy layoutmi."""
    title = next((p for p in layout.placeholders if p.placeholder_format.type in TITLE_TYPES), None)
    body = next((p for p in layout.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.OBJECT), None)
    if title is None or body is None or title.top is None or body.top is None:
        return True
    return title.top + title.height // 2 < body.top


def _pick_layouts(prs) -> dict:
    used = [s.slide_layout for s in prs.slides]
    # Ba'zi shablonlarda bir nechta master bor — dizayn namuna slaydlar ishlatgan masterda
    main = used[0].slide_master if used else prs.slide_master
    layouts = list(main.slide_layouts) + [
        l for m in prs.slide_masters if m is not main for l in m.slide_layouts
    ]

    def first(cands, pred):
        return next((l for l in cands if pred(l)), None)

    is_cover = lambda l: PP_PLACEHOLDER.CENTER_TITLE in _ph_types(l)  # noqa: E731
    is_content = lambda l: (  # noqa: E731
        PP_PLACEHOLDER.TITLE in _ph_types(l) and PP_PLACEHOLDER.OBJECT in _ph_types(l)
        and _ph_types(l).count(PP_PLACEHOLDER.OBJECT) == 1
    )
    cover = first(used, is_cover) or first(used[:1], _has_title) or first(layouts, is_cover) or layouts[0]

    content = []
    for l in used:
        if is_content(l) and _title_on_top(l) and all(l is not c for c in content):
            content.append(l)
    if not content:
        content = [first(layouts, is_content) or first(layouts, lambda l: _has_title(l) and _has_body(l))]

    closing = first(layouts, lambda l: l.name.lower() == "closing" and _has_title(l)) or cover
    return {"cover": cover, "content": content, "closing": closing}


def _capture_decor(slide) -> list:
    """Namuna slayddagi matnsiz bezaklarni (fon rasmlari, shakllar) saqlab oladi."""
    items = []
    for sh in slide.shapes:
        if sh.is_placeholder or (sh.has_text_frame and sh.text_frame.text.strip()):
            continue
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            items.append(("pic", sh.image.blob, sh.left, sh.top, sh.width, sh.height))
        elif not sh._element.xpath(".//@r:embed | .//@r:id | .//@r:link"):
            items.append(("xml", copy.deepcopy(sh._element)))
    return items


def _apply_decor(slide, items) -> None:
    tree = slide.shapes._spTree
    pos = 2  # nvGrpSpPr va grpSpPr dan keyin — ya'ni eng orqa qatlam
    for item in items:
        if item[0] == "pic":
            _, blob, left, top, w, h = item
            el = slide.shapes.add_picture(io.BytesIO(blob), left, top, w, h)._element
        else:
            el = copy.deepcopy(item[1])
        tree.remove(el) if el.getparent() is tree else None
        tree.insert(pos, el)
        pos += 1


def _remove_all_slides(prs) -> None:
    sld_ids = prs.slides._sldIdLst
    for sld_id in list(sld_ids):
        prs.part.drop_rel(sld_id.rId)
        sld_ids.remove(sld_id)


def _strip_watermarks(prs) -> None:
    """Shablon mualliflarining '@..._bot' yozuvlarini master va layoutlardan olib tashlaydi."""
    parts = list(prs.slide_masters) + [l for m in prs.slide_masters for l in m.slide_layouts]
    for part in parts:
        for shape in list(part.shapes):
            if shape.has_text_frame and "_bot" in shape.text_frame.text:
                if shape.is_placeholder:
                    shape.text_frame.text = ""  # placeholderni o'chirsak dizayn buziladi
                else:
                    shape._element.getparent().remove(shape._element)


def _title_ph(slide):
    return next((p for p in slide.placeholders if p.placeholder_format.type in TITLE_TYPES), None)


def _body_phs(slide):
    phs = [p for p in slide.placeholders if p.placeholder_format.type in BODY_TYPES]
    return sorted(phs, key=lambda p: -(p.width or 0) * (p.height or 0))


def _clean_empty(slide) -> None:
    for ph in list(slide.placeholders):
        if ph.has_text_frame and not ph.text_frame.text.strip():
            ph._element.getparent().remove(ph._element)
        elif ph.placeholder_format.type == PP_PLACEHOLDER.PICTURE:
            ph._element.getparent().remove(ph._element)


def _set_text(ph, text: str, size: int | None = None) -> None:
    tf = ph.text_frame
    tf.text = text
    tf.word_wrap = True
    if size:
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size = Pt(size)


def _set_bullets(ph, bullets: list[str], size: int) -> None:
    tf = ph.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = b
        run.font.size = Pt(size)
        p.space_after = Pt(6)


def _bullet_size(bullets: list[str], narrow: bool) -> int:
    chars = sum(len(b) for b in bullets)
    size = 20 if chars < 280 else 18 if chars < 420 else 16 if chars < 560 else 14
    return size - 2 if narrow else size


def _media_area(body, slide_w: int):
    """Rasm yoki diagramma uchun o'ng tomondagi joy (left, top, width, height).

    Matn maydoni slaydning chap qismida bo'lsa, bo'sh o'ng tomon olinadi,
    aks holda matn maydoni toraytiriladi.
    """
    left, top, width, height = body.left, body.top, body.width, body.height
    gap = Emu(Pt(18))
    margin = int(slide_w * 0.05)
    if left + width < slide_w * 0.66:
        area_left, area_w = left + width + gap, slide_w - margin - (left + width + gap)
    else:
        text_w = int(width * 0.55)
        body.left, body.top, body.width, body.height = left, top, text_w, height
        area_left, area_w = left + text_w + gap, width - text_w - gap
    return area_left, top, area_w, height


def _add_picture(slide, area, image: bytes) -> None:
    """Rasmni 4:3 nisbatda, markazdan kesib, oq ramka bilan qo'yadi."""
    from PIL import Image

    area_left, top, area_w, height = area
    pic_w, pic_h = area_w, int(area_w * 3 / 4)
    if pic_h > height:
        pic_h, pic_w = height, int(height * 4 / 3)
    pic_left = area_left + (area_w - pic_w) // 2
    pic_top = top + Emu(Pt(6))  # matn bilan bir chiziqda boshlanadi

    with Image.open(io.BytesIO(image)) as im:
        iw, ih = im.size
    pic = slide.shapes.add_picture(io.BytesIO(image), pic_left, pic_top, pic_w, pic_h)
    # Oq ramka: rasm har qanday fonda "foto-karta" kabi aniq ajralib turadi
    pic.line.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    pic.line.width = Pt(3)
    box, img = pic_w / pic_h, iw / ih
    if img > box:  # rasm kengroq — yon tomonlardan kesamiz
        cut = (1 - box / img) / 2
        pic.crop_left = pic.crop_right = cut
    else:
        cut = (1 - img / box) / 2
        pic.crop_top = pic.crop_bottom = cut


CHART_INK = RGBColor(0x1F, 0x29, 0x33)
CHART_MUTED = RGBColor(0x5B, 0x66, 0x70)
CHART_GRID = RGBColor(0xE3, 0xE7, 0xEB)


def _set_alpha(shape, percent: int) -> None:
    """Shakl to'ldirilishiga shaffoflik beradi (python-pptx'da to'g'ridan-to'g'ri API yo'q)."""
    from pptx.oxml.ns import qn

    clr = shape.fill._xPr.find(qn("a:solidFill"))[0]
    alpha = clr.makeelement(qn("a:alpha"), {"val": str(percent * 1000)})
    clr.append(alpha)


CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.DOUGHNUT,
}


def _card(slide, left, top, width, height, alpha: int = 95, radius: float = 0.06):
    """Oq yarim shaffof kartochka — infografika va diagramma har qanday fonda o'qiladi."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.adjustments[0] = radius
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    _set_alpha(card, alpha)
    card.line.fill.background()
    card.shadow.inherit = False
    return card


def _box(slide, left, top, width, height, text, size, bold=False, color=None, theme=None,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=None):
    """Oddiy matn qutisi (avtomatik kichrayadi)."""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    if font:
        r.font.name = font
    if theme is not None:
        r.font.color.theme_color = theme
    else:
        r.font.color.rgb = color or CHART_INK
    return tb


def _num_format(unit: str) -> str:
    unit = (unit or "").replace('"', "").strip()
    if not unit:
        return "General"
    return f'General"{unit}"' if unit in ("%", "‰") else f'General" {unit}"'


def _fmt_value(v, unit: str) -> str:
    v = int(v) if float(v).is_integer() else v
    text = f"{v:,}".replace(",", " ") if isinstance(v, int) else str(v)
    unit = (unit or "").strip()
    return f"{text}{unit}" if unit in ("%", "‰") else f"{text} {unit}".strip()


def _add_chart(slide, area, chart: dict, font: str | None = None) -> None:
    """Premium uslubdagi haqiqiy PowerPoint diagrammasi (raqamlarini PowerPoint'da tahrirlash mumkin).

    Ranglar shablon mavzusidan olinadi (accent), eng katta qiymat urg'u rangida ajratiladi,
    tepada asosiy raqam (insight) chiqadi — premium-engine'dagi chart_hero g'oyasi.
    """
    area_left, top, area_w, height = area
    kind = chart.get("type", "column")
    unit = chart.get("unit", "")
    values = list(chart["values"])
    card_h = min(height, int(area_w * 0.86))
    _card(slide, area_left, top, area_w, card_h)
    pad = Emu(Pt(16))
    x, y, w = area_left + pad, top + pad, area_w - 2 * pad

    # Sarlavha va asosiy raqam
    head_h = Emu(Pt(40))
    peak = max(range(len(values)), key=lambda i: values[i])
    title = chart.get("title") or ""
    if kind != "pie":
        _box(slide, x, y, int(w * 0.62), head_h, title, 13, bold=True, font=font, anchor=MSO_ANCHOR.MIDDLE)
        _box(slide, x + int(w * 0.62), y, int(w * 0.38), head_h, _fmt_value(values[peak], unit), 24, bold=True,
             theme=MSO_THEME_COLOR.ACCENT_1, align=PP_ALIGN.RIGHT, font=font, anchor=MSO_ANCHOR.MIDDLE)
        _box(slide, x + int(w * 0.62), y + head_h, int(w * 0.38), Emu(Pt(14)), chart["labels"][peak], 10,
             color=CHART_MUTED, align=PP_ALIGN.RIGHT, font=font)
    else:
        _box(slide, x, y, w, head_h, title, 13, bold=True, font=font, anchor=MSO_ANCHOR.MIDDLE)
    cy = y + head_h + Emu(Pt(14))
    ch_h = top + card_h - pad - cy

    data = CategoryChartData()
    data.categories = chart["labels"]
    data.add_series(unit or title or "", values, number_format=_num_format(unit))
    gf = slide.shapes.add_chart(CHART_TYPES.get(kind, CHART_TYPES["column"]), x, cy, w, ch_h, data)
    ch = gf.chart
    ch.has_title = False
    ch.font.size = Pt(11)
    ch.font.color.rgb = CHART_MUTED
    if font:
        ch.font.name = font
    plot = ch.plots[0]
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.font.size = Pt(11)
    dl.font.bold = True
    dl.font.color.rgb = CHART_INK
    dl.number_format = _num_format(unit)
    dl.number_format_is_linked = False
    series = plot.series[0]

    if kind == "pie":
        plot.vary_by_categories = True
        accents = [MSO_THEME_COLOR.ACCENT_1, MSO_THEME_COLOR.ACCENT_2, MSO_THEME_COLOR.ACCENT_3,
                   MSO_THEME_COLOR.ACCENT_4, MSO_THEME_COLOR.ACCENT_5, MSO_THEME_COLOR.ACCENT_6]
        for i, pt in enumerate(series.points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.theme_color = accents[i % len(accents)]
            pt.format.line.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            pt.format.line.width = Pt(1.5)
        _set_hole(plot, 58)
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.RIGHT
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(11)
        ch.legend.font.color.rgb = CHART_INK
        return

    ch.has_legend = False
    plot.vary_by_categories = False
    if kind in ("column", "bar"):
        plot.gap_width = 45
        plot.overlap = 0
        try:
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
        except Exception:
            pass
        # Eng katta qiymat — to'q urg'u, qolganlari och tusda
        for i, pt in enumerate(series.points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
            if i != peak:
                pt.format.fill.fore_color.brightness = 0.55
            pt.format.line.fill.background()
    else:  # line
        series.format.line.color.theme_color = MSO_THEME_COLOR.ACCENT_1
        series.format.line.width = Pt(3)
        series.smooth = False
        series.marker.style = XL_MARKER_STYLE.CIRCLE
        series.marker.size = 9
        series.marker.format.fill.solid()
        series.marker.format.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        series.marker.format.line.color.theme_color = MSO_THEME_COLOR.ACCENT_1
        series.marker.format.line.width = Pt(2.25)
        try:
            dl.position = XL_LABEL_POSITION.ABOVE
        except Exception:
            pass
    try:
        va, ca = ch.value_axis, ch.category_axis
        va.visible = False
        va.has_major_gridlines = kind == "line"
        if kind == "line":
            va.major_gridlines.format.line.color.rgb = CHART_GRID
        va.has_minor_gridlines = False
        ca.format.line.color.rgb = CHART_GRID
        ca.tick_labels.font.size = Pt(11)
        ca.tick_labels.font.color.rgb = CHART_MUTED
        from pptx.enum.chart import XL_TICK_MARK
        ca.major_tick_mark = XL_TICK_MARK.NONE
    except Exception:
        pass


def _set_hole(plot, percent: int) -> None:
    from pptx.oxml.ns import qn

    dn = plot._element
    hole = dn.find(qn("c:holeSize"))
    if hole is None:
        hole = dn.makeelement(qn("c:holeSize"), {})
        dn.append(hole)
    hole.set("val", str(percent))


# ---------------- Infografika ----------------

def _add_stats(slide, area, items: list, font: str | None = None) -> None:
    """Katta raqamlar (premium-engine: stat_dominant / stat_row): o'ng tomonda 2–4 ta kartochka."""
    area_left, top, area_w, height = area
    items = items[:4]
    n = len(items)
    cols = 2 if n == 4 else 1
    rows = (n + cols - 1) // cols
    gap = Emu(Pt(10))
    cw = (area_w - gap * (cols - 1)) // cols
    chh = min((height - gap * (rows - 1)) // rows, Emu(Pt(118)))
    for i, it in enumerate(items):
        r, c = divmod(i, cols)
        x = area_left + c * (cw + gap)
        y = top + r * (chh + gap)
        _card(slide, x, y, cw, chh, alpha=95, radius=0.08)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y + Emu(Pt(14)), Emu(Pt(4)), chh - Emu(Pt(28)))
        bar.fill.solid()
        bar.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
        bar.line.fill.background()
        bar.shadow.inherit = False
        px = x + Emu(Pt(18))
        pw = cw - Emu(Pt(30))
        _box(slide, px, y + Emu(Pt(12)), pw, int(chh * 0.5), str(it.get("value", "")), 30 if cols == 1 else 24,
             bold=True, theme=MSO_THEME_COLOR.ACCENT_1, font=font, anchor=MSO_ANCHOR.BOTTOM)
        _box(slide, px, y + Emu(Pt(14)) + int(chh * 0.5), pw, int(chh * 0.5) - Emu(Pt(24)),
             str(it.get("label", "")), 12, font=font)


def _add_flow(slide, box, items: list, kind: str, font: str | None = None, bullets: list | None = None) -> None:
    """Bosqichlar (process_stair) yoki vaqt chizig'i (timeline_h) — matn maydoni o'rnida, kartochkada."""
    left, top, width, height = box
    items = items[:5]
    n = len(items)
    _card(slide, left, top, width, height, alpha=95, radius=0.04)
    pad = Emu(Pt(18))
    x0, w = left + pad, width - 2 * pad
    col = w // n
    line_y = top + int(height * (0.48 if kind == "timeline" else 0.26))
    # Gorizontal chiziq
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x0 + col // 2, line_y - Emu(Pt(1)), col * (n - 1), Emu(Pt(2)))
    ln.fill.solid()
    ln.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
    ln.fill.fore_color.brightness = 0.5
    ln.line.fill.background()
    ln.shadow.inherit = False
    d = Emu(Pt(40 if kind == "steps" else 18))
    for i, it in enumerate(items):
        cx = x0 + col * i + col // 2
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - d // 2, line_y - d // 2, d, d)
        dot.fill.solid()
        dot.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
        if i and kind == "timeline":
            dot.fill.fore_color.brightness = 0.25
        dot.line.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        dot.line.width = Pt(2.5)
        dot.shadow.inherit = False
        if kind == "steps":
            tf = dot.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(i + 1)
            r.font.size = Pt(16)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        tx = x0 + col * i + Emu(Pt(6))
        tw = col - Emu(Pt(12))
        label, text = str(it.get("label", "")), str(it.get("text", ""))
        if kind == "timeline" and i % 2 == 0:  # yuqori-pastki navbat bilan, o'qish oson
            _box(slide, tx, top + pad, tw, line_y - top - pad - Emu(Pt(18)) - Emu(Pt(40)), text, 13,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM, font=font)
            _box(slide, tx, line_y - Emu(Pt(14)) - Emu(Pt(28)), tw, Emu(Pt(28)), label, 20, bold=True,
                 theme=MSO_THEME_COLOR.ACCENT_1, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM, font=font)
        else:
            ly = line_y + d // 2 + Emu(Pt(10))
            _box(slide, tx, ly, tw, Emu(Pt(28)), label, 17 if kind == "steps" else 20, bold=True,
                 theme=MSO_THEME_COLOR.ACCENT_1 if kind == "timeline" else None, align=PP_ALIGN.CENTER, font=font)
            bottom = top + int(height * 0.66) if (bullets and kind == "steps") else top + height - pad
            _box(slide, tx, ly + Emu(Pt(34)), tw, bottom - ly - Emu(Pt(34)), text, 13,
                 align=PP_ALIGN.CENTER, font=font)
    if bullets and kind == "steps":
        # Bosqichlar ostida 2–3 qisqa xulosa punkti
        sep = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x0, top + int(height * 0.69), w, Emu(Pt(1)))
        sep.fill.solid()
        sep.fill.fore_color.rgb = CHART_GRID
        sep.line.fill.background()
        sep.shadow.inherit = False
        tb = slide.shapes.add_textbox(x0, top + int(height * 0.72), w, int(height * 0.28) - pad)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        for j, b in enumerate(bullets[:3]):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = "•  " + b
            r.font.size = Pt(14)
            r.font.color.rgb = CHART_INK
            if font:
                r.font.name = font
            p.space_after = Pt(4)


def apply_fonts(prs, fonts: dict | None) -> None:
    """Mijoz tanlagan shriftlar: sarlavhalar va asosiy matn uchun alohida."""
    if not fonts or not (fonts.get("title") or fonts.get("body")):
        return
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_chart and fonts.get("body"):
                shape.chart.font.name = fonts["body"]
            if not shape.has_text_frame:
                continue
            is_title = shape.is_placeholder and shape.placeholder_format.type in TITLE_TYPES
            name = fonts.get("title") if is_title else fonts.get("body")
            if not name:
                continue
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name = name


def _title_size(ph, text: str, cover: bool) -> int:
    """Sarlavha o'lchami: matn uzunligi va eng uzun so'z maydonga sig'ishi bo'yicha."""
    n = len(text)
    if cover:
        size = 44 if n <= 28 else 40 if n <= 45 else 34 if n <= 70 else 28
    else:
        size = 32 if n <= 40 else 28 if n <= 60 else 24
    width_pt = (ph.width or 0) / 12700
    longest = max((len(w) for w in text.split()), default=1)
    if width_pt:
        size = min(size, int(width_pt / (longest * 0.58)))
    return max(size, 16)


def build(template: Path, out: Path, lang: str, outline: dict, content: dict, images: list | None = None,
          fonts: dict | None = None) -> Path:
    prs = Presentation(str(template))
    lay = _pick_layouts(prs)
    slides = list(prs.slides)
    decor_for = lambda layout, default=None: next(  # noqa: E731
        (_capture_decor(sl) for sl in slides if sl.slide_layout is layout), default or []
    )
    cover_decor = _capture_decor(slides[0]) if slides and slides[0].slide_layout is lay["cover"] else []
    content_decor = [decor_for(l) for l in lay["content"]]
    closing_decor = decor_for(lay["closing"]) if lay["closing"] is not lay["cover"] else cover_decor
    _remove_all_slides(prs)
    _strip_watermarks(prs)
    labels = LABELS.get(lang, LABELS["uz"])
    images = images or [None] * len(content["slides"])

    # 1. Muqova
    s = prs.slides.add_slide(lay["cover"])
    _apply_decor(s, cover_decor)
    if (t := _title_ph(s)) is not None:
        _set_text(t, outline["title"], _title_size(t, outline["title"], cover=True))
    if (b := _body_phs(s)) and outline.get("subtitle"):
        _set_text(b[0], outline["subtitle"])
    _clean_empty(s)

    def content_slide(i, title, bullets, image=None, chart=None, layout=None, items=None):
        k = i % len(lay["content"])
        s = prs.slides.add_slide(lay["content"][k])
        _apply_decor(s, content_decor[k])
        if (t := _title_ph(s)) is not None:
            _set_text(t, title, _title_size(t, title, cover=False))
        body = _body_phs(s)
        font = (fonts or {}).get("body")
        if body and layout in ("steps", "timeline") and items and len(items) >= 2:
            # Bosqich/vaqt chizig'i butun matn maydonini egallaydi; punktlar o'rniga infografika
            b = body[0]
            # Infografika uchun matn maydonidan kengroq joy: slayd kengligining 90%, pastgacha
            sw, sh = prs.slide_width, prs.slide_height
            left = min(b.left, int(sw * 0.05))
            box = (left, b.top, sw - 2 * left, max(b.height, int(sh * 0.94) - b.top))
            b._element.getparent().remove(b._element)
            try:
                _add_flow(s, box, items, layout, font, bullets)
            except Exception:
                log.exception("Infografika chizilmadi")
            _clean_empty(s)
            return
        media = False
        if body:
            want_stats = layout == "stats" and items and len(items) >= 2
            if chart or image or want_stats:
                area = _media_area(body[0], prs.slide_width)
                try:
                    if chart:
                        _add_chart(s, area, chart, font)
                    elif want_stats:
                        _add_stats(s, area, items, font)
                    else:
                        _add_picture(s, area, image)
                    media = True
                except Exception:  # buzilgan rasm yoki diagramma butun taqdimotni to'xtatmasin
                    log.exception("Media qo'shilmadi")
            _set_bullets(body[0], bullets, _bullet_size(bullets, narrow=media))
        _clean_empty(s)

    # 2. Reja
    content_slide(0, labels["agenda"], [f"{i}. {t}" for i, t in enumerate(outline["slides"], 1)])

    # 3. Asosiy slaydlar
    for i, sl in enumerate(content["slides"], 1):
        content_slide(i, sl["title"], sl["bullets"], images[i - 1] if i - 1 < len(images) else None,
                      sl.get("chart"), sl.get("layout"), sl.get("items"))

    # 4. Yakuniy slayd
    s = prs.slides.add_slide(lay["closing"])
    _apply_decor(s, closing_decor)
    if (t := _title_ph(s)) is not None:
        _set_text(t, labels["thanks"], _title_size(t, labels["thanks"], cover=True))
    if (b := _body_phs(s)) and content.get("closing"):
        _set_text(b[0], content["closing"])
    _clean_empty(s)

    apply_fonts(prs, fonts)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out
