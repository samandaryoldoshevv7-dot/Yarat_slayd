"""Tahrirlash: saqlangach fayl qayta yig'iladi; o'chirish; AI javobini tozalash."""
import asyncio
from io import BytesIO

import pytest
from pptx import Presentation

from app import ai, db, images, service
from conftest import auth, make_user
from test_access import make_order


def texts(data: bytes) -> str:
    prs = Presentation(BytesIO(data))
    return " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)


def test_save_rebuilds_pptx_with_new_text(client):
    token, oid = make_order(client, 30)
    p = client.get(f"/api/orders/{oid}/project", headers=auth(token)).json()
    slides = [{"title": s["title"], "bullets": s["bullets"], "chart": None, "layout": None, "items": []}
              for s in p["content"]["slides"]]
    slides[0]["title"] = "Yangi sarlavha 123"
    slides[0]["bullets"] = ["Tahrirlangan punkt"]
    r = client.post(f"/api/orders/{oid}/save", headers=auth(token),
                    json={"title": "Tahrirlangan taqdimot", "subtitle": "", "closing": "", "slides": slides,
                          "template": "02"})
    assert r.status_code == 200, r.text
    assert r.json()["version"] == 2
    data = client.get(f"/api/download/{oid}", headers=auth(token)).content
    t = texts(data)
    assert "Tahrirlangan taqdimot" in t and "Yangi sarlavha 123" in t and "Tahrirlangan punkt" in t
    assert db.balance(30) == 0  # tahrirlash bepul


def test_save_rejects_wrong_slide_count(client):
    token, oid = make_order(client, 31)
    r = client.post(f"/api/orders/{oid}/save", headers=auth(token),
                    json={"title": "x", "slides": [], "template": "01"})
    assert r.status_code == 400


def test_delete_order_removes_files(client):
    token, oid = make_order(client, 32)
    assert service.project_dir(oid).exists()
    assert client.delete(f"/api/orders/{oid}", headers=auth(token)).status_code == 200
    assert not service.project_dir(oid).exists()
    assert client.get(f"/api/download/{oid}", headers=auth(token)).status_code == 404
    assert client.get("/api/orders", headers=auth(token)).json() == []


def test_delete_account_data(client):
    token, oid = make_order(client, 33)
    assert client.delete("/api/account", headers=auth(token)).status_code == 200
    user = db.get_user(33)
    assert user["full_name"] == "" and user["username"] is None
    assert not service.project_dir(oid).exists()
    assert client.get("/api/me", headers=auth(token)).json() == {"logged_in": False}
    # Qayta ro'yxatdan o'tganda sovg'a bonusi qayta berilmaydi
    assert db.register_user(33, "u", "Qaytgan", None) is False


def test_ai_image_rate_limited(client, monkeypatch):
    token, oid = make_order(client, 34)

    async def fake(queries):
        from PIL import Image
        buf = BytesIO()
        Image.new("RGB", (40, 30), (10, 120, 60)).save(buf, "JPEG")
        return [buf.getvalue() for _ in queries]
    monkeypatch.setattr(images, "fetch_images", fake)
    codes = [client.post(f"/api/orders/{oid}/image/0/ai", headers=auth(token), json={"query": "green city"}).status_code
             for _ in range(21)]
    assert codes[:20] == [200] * 20 and codes[20] == 429


def test_clean_bullets():
    seen = set()
    out = ai.clean_bullets(["  • Birinchi punkt ", "", "Birinchi punkt.", "x" * 400, 123], seen)
    assert out[0] == "Birinchi punkt"
    assert len(out) == 3 and out[1].endswith("…") and len(out[1]) <= ai.MAX_BULLET + 1
    assert ai.clean_bullets(["birinchi punkt"], seen) == []  # boshqa slayddagi takror
    many = ai.clean_bullets([f"{k}-punkt " + "so'z " * 40 for k in range(6)])
    assert 1 <= len(many) < 6 and sum(map(len, many)) <= ai.SLIDE_TEXT_BUDGET


def test_make_content_survives_broken_ai_json(monkeypatch):
    monkeypatch.setattr(ai.config, "DEMO_MODE", False)
    outline = {"title": "T", "subtitle": "", "slides": ["A", "B", "C"]}

    async def broken(*a, **k):
        return {"slides": [{"bullets": "not a list"}, None, {"bullets": ["Yaxshi punkt"], "items": {"x": 1},
                                                              "layout": "steps", "chart_type": "pie",
                                                              "chart_labels": "ab", "chart_values": [1]}]}
    monkeypatch.setattr(ai, "_ask_json", broken)
    c = asyncio.run(ai.make_content("T", "uz", outline, charts=True))
    assert [s["bullets"] for s in c["slides"]] == [[], [], ["Yaxshi punkt"]]
    assert c["closing"] == "" and c["slides"][2]["chart"] is None and c["slides"][2]["layout"] is None

    async def empty(*a, **k):
        return {"slides": []}
    monkeypatch.setattr(ai, "_ask_json", empty)
    with pytest.raises(ai.AIError):
        asyncio.run(ai.make_content("T", "uz", outline))


def test_make_outline_survives_broken_ai_json(monkeypatch):
    monkeypatch.setattr(ai.config, "DEMO_MODE", False)

    async def broken(*a, **k):
        return {"title": None, "slides": ["  Kirish ", "", 5]}
    monkeypatch.setattr(ai, "_ask_json", broken)
    o = asyncio.run(ai.make_outline("Mavzu", "uz", 7))
    assert o["title"] == "Mavzu" and o["slides"] == ["Kirish", "5"]


@pytest.mark.skipif(not __import__("shutil").which("soffice"), reason="LibreOffice o'rnatilmagan")
def test_save_regenerates_preview_and_pdf(client, monkeypatch):
    from app import preview
    monkeypatch.setattr(preview, "available", lambda: bool(__import__("shutil").which("pdftoppm")))
    token, oid = make_order(client, 35)
    p = client.get(f"/api/orders/{oid}/project", headers=auth(token)).json()
    assert p["previews"] > 0
    old_pdf = client.get(f"/api/orders/{oid}/pdf", headers=auth(token))
    assert old_pdf.status_code == 200 and old_pdf.content[:4] == b"%PDF"
    slides = [{"title": s["title"], "bullets": s["bullets"]} for s in p["content"]["slides"]]
    r = client.post(f"/api/orders/{oid}/save", headers=auth(token),
                    json={"title": "Yangi nom", "slides": slides, "template": p["template"]})
    assert r.status_code == 200 and r.json()["previews"] > 0
    new_pdf = client.get(f"/api/orders/{oid}/pdf", headers=auth(token))
    assert new_pdf.status_code == 200 and new_pdf.content != old_pdf.content
    # Eski versiya rasmlari o'chirilgan
    assert [d.name for d in (service.project_dir(oid) / "preview").iterdir()] == ["v2"]
